from step_1_1 import OUT_DIR
from step_1_2 import run_playwright
from step_2_2 import OUT_2_2
from step_2_3 import fetch_trends_by_filter
from step_3_1 import apply_font_style, init_docx
from docx.document import Document as DocumentObject
from docx.shared import Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from pathlib import Path
import json

def add_table(doc:DocumentObject,category:str,option:str) :
    fetch_trends_by_filter(category,option)
    imgs_path : list = json.loads(OUT_2_2.read_text(encoding="utf-8"))
    n_items = len(imgs_path)
    n_cols = 5
    n_rows = n_items//n_cols + (1 if n_items % n_cols > 0 else 0)

    para = doc.add_paragraph(style="Heading 2")
    text_filter = f"{option}의 {category} 트렌드"
    apply_font_style(para.add_run(text_filter),size_pt=15,is_bold=True)

    table = doc.add_table(rows=n_rows,cols=n_cols,style="Table Grid")
    for tr in table.rows :
        for td in tr.cells :
            if len(img_path) > 0 :
                img_path = imgs_path.pop(0)
                p_cell = td.paragraphs[0]
                p_cell.add_run().add_picture(img_path,width=Cm(3))
                p_cell.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
    doc.add_paragraph("")
    p_list = doc.add_paragraph(style="List Bullet")
    text_notice = "보다 자세한 정보는 지피티에게"
    apply_font_style(p_list.add_run(text_notice),size_pt=9)

if __name__ == "__main__" :
    doc = init_docx()
    add_table(doc,"패션뷰티","여성의류")
    doc.save(OUT_DIR/f"{Path(__file__).stem}.docx")