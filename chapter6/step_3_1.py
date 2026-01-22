from datetime import datetime as dt
from pathlib import Path
from docx import Document
from docx.document import Document as DocumentObject
from docx.shared import Pt
from docx.text.run import Run
from step_1_1 import OUT_DIR

def apply_font_style(run:Run, size_pt:int=None, is_bold:bool=None) :
    if size_pt is not None : 
        run.font.size = Pt(size_pt)
    if is_bold is not None :
        run.font.bold = is_bold

def init_docx() -> DocumentObject :
    doc = Document()
    p1 = doc.add_paragraph(style="Heading 1")
    run1 = p1.add_run("쇼핑 트렌드 보고서")
    apply_font_style(run1,size_pt=25,is_bold=True)
    return doc

if __name__ == "__main__" :
    doc = init_docx()
    doc.save(OUT_DIR/f"{Path(__file__).stem}.docx")